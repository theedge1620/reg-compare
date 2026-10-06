#!/usr/bin/env python3
"""
folder_compare.py — RegDocDiff folder comparison.

Analyzes every .docx amendment document in a folder using the same logic as the
"Document Comparison" tab of index.html, and writes a dataset of every
conflicting subsection with its redline:

    regdocdiff-folder.json   load into the "Folder Comparison" tab of index.html
    regdocdiff-folder.csv    one row per redline, [-deleted-] {+inserted+} markup
    regdocdiff-folder.xlsx   same rows, with real red strikethrough / green underline
                          (written when openpyxl is installed)

Usage:
    python folder_compare.py FOLDER [-o OUTDIR] [--recursive] [--no-csv] [--no-xlsx]

The core needs only the Python standard library. openpyxl (bundled with
Anaconda) is used for the .xlsx output if it is available.

Section numbers written without a "§" are checked against the eCFR's list of
Title 10 sections (fetched once per eCFR update and cached), so quantities such
as "0.001 rem" are not mistaken for section headings. Use --no-ecfr to skip it.

Parity note: the parsing, conflict detection and redline algorithms below are a
line-for-line port of the JavaScript in index.html. Regular expressions are
written to reproduce JavaScript semantics (ASCII \\d, \\w and \\b; Unicode \\s),
so the same document yields the same units in both. If you change one, change
the other.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
import time
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

SCHEMA = "regdiff-folder/1"      # format id; kept from the RegDiff name so existing datasets still load

# ─────────────────────────────────────────────────────────────────────────────
# JavaScript-compatible regex building blocks
# ─────────────────────────────────────────────────────────────────────────────
#
# JS \s is Unicode whitespace; JS \d, \w and \b are ASCII-only. Python's
# defaults differ on each, so the classes are spelled out explicitly.

_JS_WS_CHARS = (
    "\t\n\v\f\r   "
    + "".join(chr(c) for c in range(0x2000, 0x200B))
    + "    　﻿"
)
_WS_CLASS = "\t\n\v\f\r    -     　﻿"
S = f"[{_WS_CLASS}]"          # JS \s
NS = f"[^{_WS_CLASS}]"        # JS \S
_W = "A-Za-z0-9_"
B = f"(?:(?<=[{_W}])(?![{_W}])|(?<![{_W}])(?=[{_W}]))"   # JS \b

SECNUM = r"([0-9]+[A-Za-z]?(?:\.[0-9]+[A-Za-z]*)*)"

# Numbered instruction. Working drafts use an "XX." placeholder before the
# amendments are numbered, so those count as instructions too.
INSTRUCTION_RE = re.compile(r"^([0-9]{1,3}|X{1,3}|x{2,3})\." + S + r"+(.+)$")
# A numbered line is an instruction only if it reads like one; otherwise it is a
# numbered regulation paragraph ("3. The Commission may not require …").
INSTRUCTION_LEAD_RE = re.compile(
    r"^(?:In" + S + r"+(?:§|part" + B + "|[Aa]ppendix" + B + "|[Ss]ubpart" + B + "|[Ss]ection" + B
    + "|[Tt]able" + B + "|paragraphs?" + B + "|the" + S
    + r"+(?:definition|table|heading|introductory|undesignated|entry|list|authority|section))|§"
    + "|(?:Revise|Add|Amend|Remove|Redesignate|Republish|Reserve|Designate|Transfer|Suspend|Correct|Effective)"
    + B + "|The authority citation)")
INSTRUCTION_PASSIVE_RE = re.compile(
    B + "(?:is|are)" + S + "+(?:hereby" + S + "+)?(?:amended|revised|added|removed|redesignated|republished"
    "|reserved|corrected|transferred|suspended)" + B + "|" + B + "to read as follows" + B
    + "|" + B + "continues to read" + B, re.I)
# Unnumbered instructions (Word list numbering isn't stored as text) must be
# unmistakable: an instruction opening, amendatory wording and a § reference.
UNNUMBERED_LEAD_RE = re.compile(
    r"^(?:In" + S + "+§|Revise" + B + "|Add" + B + "|Amend" + B + "|Remove" + B + "|Redesignate" + B
    + "|Republish" + B + "|The authority citation for part" + B + ")")
AUTHORITY_LEAD_RE = re.compile(r"^The authority citation for part" + B)
AMENDATORY_RE = re.compile(
    ":" + S + r"*\Z|" + B + "to read as follows" + B + "|" + B + "in its place" + B
    + "|" + B + "remove and reserve" + B + "|" + B + "continues to read" + B
    + "|" + B + "(?:is|are)" + S + "+(?:amended|revised|added|removed|redesignated|republished|reserved)" + B
    + "|" + B + "remove the (?:words?|phrases?|references?|definitions?|sentences?|entry|entries)" + B
    + "|" + B + "redesignat", re.I)
# Lettered clauses inside one instruction: "In § 110.45: a. In paragraph …; and b. …"
CLAUSE_MARK_RE = re.compile(r"([:;.])" + S + "+(?:and" + S + r"+)?([a-z])\." + S + "+")
# A clause that reprints regulation text rather than editing it in place.
PRINTS_TEXT_RE = re.compile(r"^(?:Revise|Add|Republish)" + B + "|" + B + "to read as follows" + B, re.I)
SUBINSTR_RE = re.compile(r"^([a-z])\." + S + r"+(.+)$")
SECTION_HEAD_RE = re.compile(r"^§{1,2}" + S + "*" + SECNUM + S + "*[.:—–-]?" + S + r"*(.*)$")
DOTTED_HEAD_RE = re.compile(r"^([0-9]+(?:\.[0-9]+)+[A-Za-z]?)\.?" + S + r"+(.+)$")
# "* * *" (or longer) alone on a line — the Federal Register elision marker.
ELISION_RE = re.compile(r"^\*(?:" + S + r"*\*){2,}" + S + "*$")
SECTION_REF_RE = re.compile(
    r"(?:§{1,2}|" + B + "Sections?" + B + "|" + B + r"Secs?\.)" + S + "*" + SECNUM, re.I)
PARA_KW_RE = re.compile(B + "paragraphs?" + B, re.I)
PARA_SEP_RE = re.compile(r"^(?:" + S + "*(?:,|and|or)?" + S + "*)", re.I)
DESIGNATOR_TOKEN_RE = re.compile(r"^[A-Za-z0-9]{1,4}$")
BRACKET_ONLY_RE = re.compile(r"^\[[^\]]*\]$")
ROMAN_RE = re.compile(r"^[ivxlcdm]+$")
WS_RUN_RE = re.compile(S + "+")
DIFF_TOKEN_RE = re.compile("(" + NS + "+|" + S + "+)")
SIM_TOKEN_RE = re.compile(r"[A-Za-z0-9_'-]+")

ROMAN_SUCCESSOR = {"i": "ii", "v": "vi", "x": "xi"}


def js_trim(s: str) -> str:
    return s.strip(_JS_WS_CHARS)


def js_len(s: str) -> int:
    """String length as JavaScript counts it (UTF-16 code units)."""
    return len(s.encode("utf-16-le")) // 2


# ─────────────────────────────────────────────────────────────────────────────
# DOCX → ordered plain-text paragraphs
# ─────────────────────────────────────────────────────────────────────────────
#
# Only <w:t> runs contribute text, so tracked deletions (<w:delText>) are
# excluded and tracked insertions included — the document as it reads with all
# changes accepted.

_WORD_NS = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "http://purl.oclc.org/ooxml/wordprocessingml/main",          # Strict OOXML
)
_P_TAGS = {f"{{{ns}}}p" for ns in _WORD_NS}


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _paragraph_text(p: ET.Element) -> str:
    out: list[str] = []

    def walk(node: ET.Element) -> None:
        for child in node:
            name = _local(child.tag) if isinstance(child.tag, str) else ""
            if name == "t":
                out.append("".join(child.itertext()))
            elif name == "delText":
                continue                       # tracked deletion — omit
            elif name in ("tab", "br", "cr"):
                out.append(" ")
            else:
                walk(child)

    walk(p)
    text = "".join(out).replace(" ", " ")
    return js_trim(WS_RUN_RE.sub(" ", text))


def read_docx_paragraphs(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as z:
            xml_bytes = z.read("word/document.xml")
    except zipfile.BadZipFile:
        raise ValueError("Not a valid .docx file (no ZIP directory found).")
    except KeyError:
        raise ValueError("Malformed .docx — word/document.xml not found in the archive.")
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        raise ValueError("Could not parse the Word document body (document.xml is malformed).")

    paras = []
    for el in root.iter():
        if el.tag in _P_TAGS:
            t = _paragraph_text(el)
            if t:
                paras.append(t)
    return paras


# ─────────────────────────────────────────────────────────────────────────────
# Amendment document structure — port of parseAmendmentDoc() and helpers
# ─────────────────────────────────────────────────────────────────────────────

BARE_HEADING_MIN_SIMILARITY = 0.5


def is_bare_section_heading(number: str, heading: str, instr, title10) -> bool:
    """A bare number opening a line ("2.326 Motions to reopen.") is a section
    heading only when vouched for; otherwise it is a quantity that happens to
    start the line ("0.001 rem in any one hour"). It counts when the governing
    instruction names that section (covers sections new in the rule, not yet in
    the eCFR), or when it is a Title 10 section in the eCFR AND the text after
    it matches that section's eCFR heading — "1.5" alone is a real section."""
    if instr and number in instr["refs"]:
        return True
    if not title10 or number not in title10:
        return False
    return similarity(heading, title10[number]) >= BARE_HEADING_MIN_SIMILARITY


def read_instruction(line: str):
    """{num, body} when the line is an amendment instruction, else None. `body`
    is the instruction without its number; unnumbered instructions have num ''."""
    m = INSTRUCTION_RE.match(line)
    if m and not DOTTED_HEAD_RE.match(line):
        body = m.group(2)
        if INSTRUCTION_LEAD_RE.match(body) or INSTRUCTION_PASSIVE_RE.search(body):
            return {"num": m.group(1), "body": body}
        return None
    if UNNUMBERED_LEAD_RE.match(line) and AMENDATORY_RE.search(line) \
            and (instruction_section_refs(line) or AUTHORITY_LEAD_RE.match(line)):
        return {"num": "", "body": line}
    return None


def split_instruction_clauses(body: str):
    """"In § 110.45: a. In paragraph (b)(3) …; and b. …" → (lead, [clauses])."""
    marks, expect = [], "a"
    for m in CLAUSE_MARK_RE.finditer(body):
        if m.group(2) != expect:
            continue
        marks.append(m)
        expect = chr(ord(expect) + 1)
    if not marks:
        return body, []
    lead = js_trim(body[:marks[0].start() + 1])
    clauses = [js_trim(body[m.end():marks[i + 1].start() if i + 1 < len(marks) else len(body)])
               for i, m in enumerate(marks)]
    return lead, clauses


# Quoted phrases in an instruction are text being removed or inserted
# ('remove the references “§ 50.83 or”'), not what the instruction amends.
QUOTED_RE = re.compile(r'[“"][^“”"]*[”"]|``[^`]*?' + "''")


def unquoted(text: str) -> str:
    return QUOTED_RE.sub(" ", text)


def instruction_section_refs(text: str) -> list[str]:
    refs: list[str] = []
    for m in SECTION_REF_RE.finditer(unquoted(text)):
        number = m.group(1).rstrip(".")
        if number not in refs:
            refs.append(number)
    return refs


# ─────────────────────────────────────────────────────────────────────────────
# eCFR Title 10 section list — port of getTitle10Sections()
# ─────────────────────────────────────────────────────────────────────────────
#
#   GET /api/versioner/v1/titles.json                     → current date
#   GET /api/versioner/v1/structure/{date}/title-10.json  → every section's
#                                                            identifier + heading
#
# Cached per eCFR date; if the eCFR can't be reached the newest cached list is
# used, and failing that none (then only instructions can vouch for a heading).

ECFR_API = "https://www.ecfr.gov/api/versioner/v1"
SECTION_ID_RE = re.compile(r"^[0-9]+[A-Za-z]?(?:\.[0-9]+[A-Za-z]*)*$")   # skips "1.40-1.41"


def default_cache_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA") or os.environ.get("XDG_CACHE_HOME") or str(Path.home() / ".cache")
    return Path(base) / "regdiff"


def _get_json(url: str, timeout: int = 60):
    req = urllib.request.Request(url, headers={"User-Agent": "regdocdiff-folder-compare/1",
                                               "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.load(r)


def title10_sections_from_structure(root: dict) -> dict[str, str]:
    sections: dict[str, str] = {}
    stack = [root]
    while stack:
        node = stack.pop()
        if node.get("type") == "section" and SECTION_ID_RE.match(node.get("identifier") or ""):
            sections[node["identifier"]] = node.get("label_description") or ""
        stack.extend(node.get("children") or [])
    return sections


def load_title10_sections(cache_dir: Path | None = None, log=print):
    """Returns (sections, info): sections maps identifier → eCFR heading, or is
    None when no list could be had; info is {date, sections, source} or None."""
    cache_dir = cache_dir or default_cache_dir()

    def read(path: Path):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) and data else None
        except (OSError, ValueError):
            return None

    try:
        titles = _get_json(f"{ECFR_API}/titles.json")
        date = next(t["up_to_date_as_of"] for t in titles["titles"] if t["number"] == 10)
        path = cache_dir / f"title10-sections-{date}.json"
        sections = read(path) if path.exists() else None
        source = "cache"
        if sections is None:
            sections = title10_sections_from_structure(
                _get_json(f"{ECFR_API}/structure/{date}/title-10.json"))
            source = "eCFR"
            try:
                cache_dir.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps(sections, ensure_ascii=False), encoding="utf-8")
            except OSError as e:
                log(f"note: could not cache the eCFR section list ({e})")
        return sections, {"date": date, "sections": len(sections), "source": source}
    except Exception as e:                                # network, HTTP or JSON failure
        cached = sorted(cache_dir.glob("title10-sections-*.json")) if cache_dir.is_dir() else []
        for path in reversed(cached):                     # dates sort lexically: newest last
            sections = read(path)
            if sections:
                date = path.stem.removeprefix("title10-sections-")
                log(f"warning: eCFR unreachable ({e}); using cached section list from {date}")
                return sections, {"date": date, "sections": len(sections), "source": "stale cache"}
        log(f"warning: eCFR unreachable ({e}) and no cached section list — an unmarked "
            f"number will count as a section only when an instruction names it")
        return None, None


def designator_level(token: str, stack: list[str], next_token: str | None) -> int:
    """(a) letter → (1) digit → (i) roman → (A) capital. A lone i/v/x is
    resolved from the next designator: a letter's children are digits, a
    roman numeral's children are capitals."""
    if re.fullmatch(r"[0-9]+", token):
        return 2
    if re.fullmatch(r"[A-Z]+", token):
        return 4
    if re.fullmatch(r"[a-z]", token):
        if token not in "ivx":
            return 1
        if next_token:
            if re.fullmatch(r"[0-9]+", next_token):
                return 1
            if re.fullmatch(r"[A-Z]+", next_token):
                return 3
            if next_token == ROMAN_SUCCESSOR[token]:
                return 3
            if next_token == chr(ord(token) + 1):
                return 1
        prev = stack[0] if stack else ""
        if prev and re.fullmatch(r"\([a-z]\)", prev):
            if token == chr(ord(prev[1]) + 1):
                return 1
        return 3 if (len(stack) > 1 and stack[1]) else 1
    if ROMAN_RE.match(token):
        return 3
    if re.fullmatch(r"[a-z]{2,}", token):
        return 1
    return 1


def leading_designators(line: str):
    """Leading "(x)" groups → (tokens, rest), or None."""
    tokens: list[str] = []
    pos = 0
    while pos < len(line) and line[pos] == "(":
        close = line.find(")", pos)
        if close == -1:
            break
        token = line[pos + 1:close]
        if not DESIGNATOR_TOKEN_RE.match(token):
            break
        tokens.append(token)
        pos = close + 1
    if not tokens:
        return None
    return tokens, js_trim(line[pos:])


def resolve_path(tokens: list[str], stack: list[str], next_token: str | None) -> str:
    """A compound designator hangs off the level of its first token, so
    "(5)(i)" under "(c)(2)(ii)" resolves to "(c)(5)(i)"."""
    level = designator_level(tokens[0], stack, tokens[1] if len(tokens) > 1 else next_token)
    keep = max(0, level - 1)
    if len(stack) > keep:
        del stack[keep:]
    else:
        stack.extend([""] * (keep - len(stack)))   # JS array holes join as ""
    stack.extend(f"({t})" for t in tokens)
    return "".join(stack)


def instruction_para_targets(text: str) -> list[str]:
    targets: list[str] = []
    text = unquoted(text)
    for m in PARA_KW_RE.finditer(text):
        pos = m.end()
        while True:
            sep = PARA_SEP_RE.match(text[pos:])
            start = pos + (len(sep.group(0)) if sep else 0)
            des = leading_designators(text[start:])
            if not des:
                break
            path = "".join(f"({t})" for t in des[0])
            if path not in targets:
                targets.append(path)
            pos = start + len(path)
    return targets


# ── non-regulation text — port of the JS constants of the same names ────────
#
# Structural and editorial lines that are not regulation text. Left alone,
# each is appended to whatever paragraph precedes it and makes that paragraph
# differ between documents.

PART_HEAD_RE = re.compile(r"^PART" + S + r"+[0-9]+[A-Za-z]?" + S + "*[—–-]")
SUBPART_HEAD_RE = re.compile(r"^Subpart" + S + r"+[A-Z]{1,3}(?:" + S + "*[—–-]|" + S + "+[A-Z][a-z])")
APPENDIX_HEAD_RE = re.compile(r"^Appendix" + S + r"+[A-Z0-9]{1,3}" + S + "+to" + S + "+Part"
                              + S + r"+[0-9]+" + S + "*[—–-]", re.I)
EDITORIAL_RE = re.compile(r"^(?:Authority:|Source:|Editorial Note:|Secs?\.\Z|_{3,}\Z)")
TOC_SEC_RE = re.compile(r"^Secs?\.\Z")
FOOTNOTE_TEXT_RE = re.compile(r"^\[([0-9]{1,3})\]" + S)
BARE_FOOTNOTE_RE = re.compile(r"^([0-9]{1,2})" + S + "+" + NS)
FOOTNOTE_MARK_RE = re.compile(r"([A-Za-z)\]”\"’]|[A-Za-z][.,;:])([0-9]{1,2})(?=" + S + r"|\Z)")
BRACKET_MARK_RE = re.compile(r"\[([0-9]{1,3})\]")
ISSUANCE_RE = re.compile(r"^For the reasons set (?:out|forth) in the preamble", re.I)
SIGNATURE_RE = re.compile(r"^(?:Dated" + B + r"|\[FR Doc\.)")


def regulation_span(paras: list[str]) -> tuple[int, int]:
    """[begin, end) of the paragraphs holding regulation text: after the words
    of issuance when present, and before the signature block."""
    begin = next((i + 1 for i, p in enumerate(paras) if ISSUANCE_RE.match(js_trim(p))), 0)
    for i in range(begin, len(paras)):
        if SIGNATURE_RE.match(js_trim(paras[i])):
            return begin, i
    return begin, len(paras)


# ── definitions — port of definitionTerm(), termKey(), instructionDefinitionTerms()

DEFINITION_RE = re.compile(
    r"^(?:(?:For (?:the )?purposes of|As used in) this (?:part|section|subpart|chapter)[,:—–-]?" + S + "+)?"
    + r"(?:The terms?" + S + r'+)?[“"]?([A-Z0-9](?:[^“”".;:]|\.(?=' + NS + r')){0,100}?)[”"]?,?' + S
    + r"+(?:means|mean|has the (?:same )?meaning|is defined as|refers to|includes)" + B)
DEFINITION_PERIOD_RE = re.compile(r'^([A-Z][^“”".;:]{0,80}?)\.(?:' + S + r"+(?=[(A-Z])|\Z)")
TERM_QUALIFIER_RE = re.compile(r"," + S + "*(?:as used in|for (?:the )?purposes of)" + B + r".*\Z", re.I | re.S)
NOT_A_TERM_RE = re.compile(B + "(?:shall|must|may|will|should|would|can|could)" + B, re.I)
DEF_TERMS_RE = re.compile(B + "definitions?" + S + "+(?:of|for)" + S + "+(?:the" + S + "+terms?" + S + "+)?", re.I)
DEF_TERM_BARE_RE = re.compile(
    r'^([A-Za-z][^,.;:“”"]{0,80}?)(?=' + S + "+(?:to read|in alphabetical order|is|are|and add|and remove|remove|from)"
    + B + r'|[,.;:]|\Z)')
DEF_TERM_QUOTE_RE = re.compile(r"^(?:" + S + "*(?:,|;|and)" + S + "*)*" + S + "*[“\"]([^”\"]+)[”\"]", re.I)
TERM_TAIL_RE = re.compile(r"[,.;:" + _WS_CLASS + r"]+\Z")


def definition_term(line: str) -> str | None:
    """"Term means …", or the CFR's other style "Term. (1) When applied …"."""
    m, max_words = DEFINITION_RE.match(line), 12
    if not m:
        m, max_words = DEFINITION_PERIOD_RE.match(line), 8
    if not m:
        return None
    term = js_trim(TERM_QUALIFIER_RE.sub("", m.group(1)))
    if not term or len(WS_RUN_RE.split(term)) > max_words or NOT_A_TERM_RE.search(term):
        return None
    return term


def term_key(term: str) -> str:
    """Terms match regardless of case or quote style."""
    return norm_for_compare(TERM_TAIL_RE.sub("", term))


def instruction_definition_terms(text: str) -> list[str]:
    terms: list[str] = []
    for m in DEF_TERMS_RE.finditer(text):
        pos = m.end()
        quoted = 0
        while True:
            q = DEF_TERM_QUOTE_RE.match(text[pos:])
            if not q:
                break
            term = js_trim(TERM_TAIL_RE.sub("", q.group(1)))
            if term and term not in terms:
                terms.append(term)
            pos += len(q.group(0))
            quoted += 1
        if not quoted:
            # Unquoted (italics lost): "the definition for Sealed Source … to read as follows"
            b = DEF_TERM_BARE_RE.match(text[pos:])
            term = js_trim(b.group(1)) if b else ""
            if term and len(WS_RUN_RE.split(term)) <= 8 and term not in terms:
                terms.append(term)
    return terms


def substantive_key(s: str) -> str:
    """Comparison key ignoring formatting: case, spacing, punctuation, quote
    and dash styles, hyphenation and footnote markers. Port of substantiveKey()."""
    s = BRACKET_MARK_RE.sub(" ", s)
    s = FOOTNOTE_MARK_RE.sub(r"\1", s)
    n = len(s)
    out = []
    for i, ch in enumerate(s):
        if ch in "-‐‑" and 0 < i < n - 1 and s[i - 1].isalnum() and s[i + 1].isalnum():
            continue                                   # "non-power" ≡ "nonpower"
        out.append(ch if (ch.isalnum() or ch == "§") else " ")
    return re.sub(" +", " ", "".join(out)).strip(" ").lower()


def parse_amendment_doc(paras: list[str], title10: dict | None = None):
    """Returns (units, instructions). `title10` is the eCFR section list
    (identifier → heading) used to vet headings that lack a "§" and to
    recognise definitions sections.

    units:        dict "section|path" → {section, sectionTitle, path, pathLabel,
                  kind, text, hasElision, instrNums}; path is the matching key,
                  pathLabel how it reads (“focused proceeding” / “Focused Proceeding”)
    instructions: [{num, text, refs, producedUnits}]
    """
    units: dict[str, dict] = {}
    instructions: list[dict] = []
    begin, end = regulation_span(paras)
    st = {"instr": None, "section": None, "stack": [], "unit": None, "def": None}
    marks: set[str] = set()       # footnote numbers referenced in the current section

    def set_section(number: str, title: str) -> None:
        known = title or (title10.get(number) if title10 else None) or ""
        st["section"] = {"number": number, "title": title, "defs": bool(re.search("definition", known, re.I))}
        st["stack"] = []
        st["def"] = None
        marks.clear()

    def note_marks(text: str) -> None:
        for m in FOOTNOTE_MARK_RE.finditer(text):
            marks.add(m.group(2))
        for m in BRACKET_MARK_RE.finditer(text):
            marks.add(m.group(1))

    def is_footnote_text(line: str) -> bool:
        if FOOTNOTE_TEXT_RE.match(line):
            return True
        m = BARE_FOOTNOTE_RE.match(line)
        return bool(m and m.group(1) in marks)

    def starts_section_block(i: int) -> bool:
        """A "Subpart X—Title" line is a heading only when a section heading
        (or table of contents) follows; otherwise it is regulation text."""
        for j in range(i + 1, end):
            l = js_trim(paras[j])
            if not l:
                continue
            return bool(SECTION_HEAD_RE.match(l) or TOC_SEC_RE.match(l)
                        or PART_HEAD_RE.match(l) or APPENDIX_HEAD_RE.match(l))
        return False

    def store(u: dict) -> bool:
        if not u["section"]:
            return False
        key = f"{u['section']}|{u['path']}"
        prev = units.get(key)
        if prev is not None \
                and not (u["kind"] == "text" and prev["kind"] == "instruction") \
                and not (u["kind"] == prev["kind"] and js_len(u["text"]) > js_len(prev["text"])):
            return False
        units[key] = u
        return True

    def flush() -> None:
        unit = st["unit"]
        if not unit:
            return
        text = js_trim("\n".join(unit.pop("lines")))
        if text:
            unit["text"] = text
            if store(unit) and st["instr"]:
                st["instr"]["producedUnits"] = True
        st["unit"] = None

    def open_unit(path: str, label: str | None = None) -> None:
        flush()
        sec = st["section"]
        number = sec["number"] if sec else ""
        ins = st["instr"]
        st["unit"] = {
            "section": number,
            "sectionTitle": sec["title"] if sec else "",
            "path": path,
            "pathLabel": path if label is None else label,
            "kind": "text",
            "lines": [],
            "hasElision": False,
            "instrNums": [ins["num"]] if ins and (not ins["refs"] or number in ins["refs"]) else [],
        }

    def close_instruction() -> None:
        """An instruction with no regulation text is itself the amendment; one
        that reprints text may still carry in-place edits for other paragraphs.
        Compared text omits the instruction number."""
        flush()
        ins = st["instr"]
        if not ins:
            return
        lead, clauses = split_instruction_clauses(ins["body"])
        if clauses:
            pieces = [(js_trim(f"{lead} {cl}"), cl) for cl in clauses
                      if not ins["producedUnits"] or not PRINTS_TEXT_RE.search(cl)]
        else:
            pieces = [] if ins["producedUnits"] else [(ins["body"], ins["body"])]

        for text, scope in pieces:
            sub_paths = instruction_para_targets(scope) or [""]
            # "add paragraph (1)(iii) in the definition for “Construction”"
            # targets a paragraph *within* that definition.
            terms = instruction_definition_terms(scope)
            if terms:
                paths = [(f"“{term_key(t)}”{p}", f"“{t}”{p}") for t in terms for p in sub_paths]
            else:
                paths = [(p, p) for p in sub_paths]
            for ref in ins["refs"]:
                for path, label in paths:
                    store({
                        "section": ref, "sectionTitle": "", "path": path, "pathLabel": label,
                        "kind": "instruction", "text": text,
                        "hasElision": False, "instrNums": [ins["num"]],
                    })

    def peek_next_designator(i: int):
        for j in range(i + 1, end):
            l = js_trim(paras[j])
            if not l or ELISION_RE.match(l):
                continue
            if INSTRUCTION_RE.match(l) or SECTION_HEAD_RE.match(l):
                return None
            d = leading_designators(l)
            return d[0][0] if d else None
        return None

    for idx in range(begin, end):
        line = js_trim(paras[idx])
        if not line:
            continue

        # ── elision marker ──
        if ELISION_RE.match(line):
            if st["unit"]:
                st["unit"]["hasElision"] = True
            continue

        # ── structural headings: end the section; not regulation text ──
        if PART_HEAD_RE.match(line) or APPENDIX_HEAD_RE.match(line) \
                or (SUBPART_HEAD_RE.match(line) and starts_section_block(idx)):
            flush()
            st["section"] = None
            st["stack"] = []
            st["def"] = None
            continue

        # ── editorial lines and footnotes: skipped ──
        if EDITORIAL_RE.match(line) or is_footnote_text(line):
            continue

        # ── amendment instruction ──
        instr = read_instruction(line)
        if instr:
            close_instruction()
            ins = {"num": instr["num"], "text": line, "body": instr["body"],
                   "refs": instruction_section_refs(line), "producedUnits": False}
            st["instr"] = ins
            instructions.append(ins)
            if ins["refs"]:
                set_section(ins["refs"][0], "")
            else:
                st["stack"] = []
                st["def"] = None
            continue

        # ── lettered sub-instruction ──
        if st["instr"] and not st["unit"] and SUBINSTR_RE.match(line):
            ins = st["instr"]
            ins["text"] += " " + line
            ins["body"] += " " + line
            for r in instruction_section_refs(line):
                if r not in ins["refs"]:
                    ins["refs"].append(r)
            continue

        # ── section heading ──
        # "§ 2.309 …" is explicit; an unmarked "2.326 …" must be vouched for.
        m_sec = SECTION_HEAD_RE.match(line)
        if not m_sec:
            m_bare = DOTTED_HEAD_RE.match(line)
            if m_bare and is_bare_section_heading(m_bare.group(1).rstrip("."), js_trim(m_bare.group(2)),
                                                  st["instr"], title10):
                m_sec = m_bare
        if m_sec:
            flush()
            heading = js_trim(m_sec.group(2) or "")
            # "§ 2.813 [Amended]" is a status marker, not a title.
            set_section(m_sec.group(1).rstrip("."), "" if BRACKET_ONLY_RE.match(heading) else heading)
            open_unit("")
            continue

        # ── subsection designator ──
        des = leading_designators(line)
        if des:
            tokens, rest = des
            nxt = peek_next_designator(idx)
            d = st["def"]
            # Inside a definition, deeper designators are its sub-paragraphs;
            # one at or above the definition's own level ends it.
            if d and designator_level(tokens[0], st["stack"], tokens[1] if len(tokens) > 1 else nxt) <= d["base"]:
                st["def"] = d = None
            full = resolve_path(tokens, st["stack"], nxt)
            if d:
                sub = "".join(st["stack"][d["base"]:])
                open_unit(d["path"] + sub, d["label"] + sub)
            else:
                open_unit(full)
            if ELISION_RE.match(rest):
                st["unit"]["hasElision"] = True
            elif rest:
                st["unit"]["lines"].append(rest)
                note_marks(rest)
            continue

        # ── a definition: "Term means …" in a definitions section ──
        term = definition_term(line) if st["section"] and st["section"]["defs"] else None
        if term:
            d = st["def"]
            # Definitions are top-level, or inside the single lettered paragraph
            # introducing them; anything deeper is a previous definition's sub-paragraphs.
            stack = st["stack"]
            base = d["base"] if d else (1 if len(stack) == 1 and re.fullmatch(r"\([a-z]+\)", stack[0]) else 0)
            del stack[base:]
            prefix = "".join(st["stack"])
            st["def"] = {"base": base, "path": f"{prefix}“{term_key(term)}”", "label": f"{prefix}“{term}”"}
            open_unit(st["def"]["path"], st["def"]["label"])
            st["unit"]["lines"].append(line)
            note_marks(line)
            continue

        # ── continuation ──
        unit = st["unit"]
        if unit:
            unit["lines"].append(line)
            note_marks(line)
            ins = st["instr"]
            if ins and ins["num"] not in unit["instrNums"] \
                    and (not ins["refs"] or unit["section"] in ins["refs"]):
                unit["instrNums"].append(ins["num"])

    close_instruction()
    return units, instructions


# ─────────────────────────────────────────────────────────────────────────────
# Comparison helpers — port of normForCompare(), similarity()
# ─────────────────────────────────────────────────────────────────────────────

_SQUOTE_RE = re.compile("[‘’ʼ]")
_DQUOTE_RE = re.compile("[“”„]")
_DASH_RE = re.compile("[‐-―]")


def norm_for_compare(s: str) -> str:
    s = _SQUOTE_RE.sub("'", s)
    s = _DQUOTE_RE.sub('"', s)
    s = _DASH_RE.sub("-", s)
    s = s.replace(" ", " ")
    s = WS_RUN_RE.sub(" ", s)
    return js_trim(s).lower()


def similarity(a: str, b: str) -> float:
    ta = SIM_TOKEN_RE.findall(norm_for_compare(a))
    tb = SIM_TOKEN_RE.findall(norm_for_compare(b))
    if not ta and not tb:
        return 1.0
    if not ta or not tb:
        return 0.0
    counts: dict[str, int] = {}
    for t in ta:
        counts[t] = counts.get(t, 0) + 1
    common = 0
    for t in tb:
        c = counts.get(t, 0)
        if c:
            common += 1
            counts[t] = c - 1
    return (2 * common) / (len(ta) + len(tb))


def _section_sort_key(number: str):
    key = []
    for seg in number.split("."):
        m = re.match(r"([0-9]*)(.*)", seg)
        key.append((int(m.group(1)) if m.group(1) else -1, m.group(2)))
    return key


# ─────────────────────────────────────────────────────────────────────────────
# Redline — port of computeDiff(), lcsOps(), mergeReplacements(),
# diffParagraphBlocks(); emits structured ops instead of HTML
# ─────────────────────────────────────────────────────────────────────────────
#
# A row is either {"t": "omit", "n": count} for a run of unchanged paragraphs,
# or {"t": "diff", "ops": [[op, text], ...]} where op is "=", "-" or "+".
# The old side of a row is its "=" and "-" runs; the new side its "=" and "+".

MAX_WORD_TOKENS = 2000      # same guard as computeDiff()
MAX_PARAGRAPHS = 400        # same guard as diffParagraphBlocks()


def _lcs_walk(a: list, b: list):
    """Yield ("=", i, j) / ("+", None, j) / ("-", i, None) exactly as the JS
    walks its LCS table. A shared prefix is peeled off first: the JS walk
    matches equal heads greedily, so this changes nothing but the cost."""
    p = 0
    while p < len(a) and p < len(b) and a[p] == b[p]:
        yield ("=", p, p)
        p += 1
    a2, b2 = a[p:], b[p:]
    m, n = len(a2), len(b2)
    dp = [[0] * (n + 1) for _ in range(m + 1)]
    for i in range(m - 1, -1, -1):
        ai, row, nxt = a2[i], dp[i], dp[i + 1]
        for j in range(n - 1, -1, -1):
            if ai == b2[j]:
                row[j] = nxt[j + 1] + 1
            else:
                x, y = nxt[j], row[j + 1]
                row[j] = x if x > y else y
    i = j = 0
    while i < m or j < n:
        if i < m and j < n and a2[i] == b2[j]:
            yield ("=", p + i, p + j)
            i += 1
            j += 1
        elif j < n and (i >= m or dp[i][j + 1] >= dp[i + 1][j]):
            yield ("+", None, p + j)
            j += 1
        else:
            yield ("-", p + i, None)
            i += 1


def _compress(ops):
    """Merge adjacent same-op tokens, and within each changed block put the
    deletions before the insertions so an inline redline reads "[-old-]{+new+}".
    Reordering inside a block leaves both sides of the two-column view intact:
    each side keeps its own tokens in their original order."""
    ordered: list[tuple[str, str]] = []
    block: list[tuple[str, str]] = []
    for op, text in ops:
        if op == "=":
            ordered += [x for x in block if x[0] == "-"] + [x for x in block if x[0] == "+"]
            block = []
            ordered.append((op, text))
        else:
            block.append((op, text))
    ordered += [x for x in block if x[0] == "-"] + [x for x in block if x[0] == "+"]

    out: list[list[str]] = []
    for op, text in ordered:
        if out and out[-1][0] == op:
            out[-1][1] += text
        else:
            out.append([op, text])
    return out


def word_diff(old: str, new: str):
    """Returns (ops, coarse). coarse=True when the text is too long for a
    word-level diff and the whole passage is marked deleted/inserted."""
    a = DIFF_TOKEN_RE.findall(old)
    b = DIFF_TOKEN_RE.findall(new)
    if len(a) > MAX_WORD_TOKENS or len(b) > MAX_WORD_TOKENS:
        return [["-", old], ["+", new]], True
    ops = []
    for op, i, j in _lcs_walk(a, b):
        ops.append((op, a[i] if op != "+" else b[j]))
    return _compress(ops), False


def diff_paragraph_blocks(a_text: str, b_text: str) -> list[dict]:
    a = [s for s in (js_trim(x) for x in a_text.split("\n")) if s]
    b = [s for s in (js_trim(x) for x in b_text.split("\n")) if s]

    if len(a) > MAX_PARAGRAPHS or len(b) > MAX_PARAGRAPHS:
        ops, coarse = word_diff(a_text, b_text)
        row = {"t": "diff", "ops": ops}
        if coarse:
            row["coarse"] = True
        return [row]

    ka = [substantive_key(x) for x in a]      # formatting-only differences align as unchanged
    kb = [substantive_key(x) for x in b]
    raw = [(op, a[i] if i is not None else None, b[j] if j is not None else None)
           for op, i, j in _lcs_walk(ka, kb)]

    # mergeReplacements(): pair each contiguous changed block into rows.
    merged = []
    k = 0
    while k < len(raw):
        if raw[k][0] == "=":
            merged.append(("eq", None, None))
            k += 1
            continue
        dels, inss = [], []
        while k < len(raw) and raw[k][0] != "=":
            if raw[k][0] == "-":
                dels.append(raw[k][1])
            else:
                inss.append(raw[k][2])
            k += 1
        paired = min(len(dels), len(inss))
        merged += [("rep", dels[x], inss[x]) for x in range(paired)]
        merged += [("del", d, None) for d in dels[paired:]]
        merged += [("ins", None, s) for s in inss[paired:]]

    rows: list[dict] = []
    eq_run = 0
    for kind, old, new in merged:
        if kind == "eq":
            eq_run += 1
            continue
        if eq_run:
            rows.append({"t": "omit", "n": eq_run})
            eq_run = 0
        if kind == "rep":
            ops, coarse = word_diff(old, new)
            row = {"t": "diff", "ops": ops}
            if coarse:
                row["coarse"] = True
            rows.append(row)
        elif kind == "del":
            rows.append({"t": "diff", "ops": [["-", old]]})
        else:
            rows.append({"t": "diff", "ops": [["+", new]]})
    if eq_run:
        rows.append({"t": "omit", "n": eq_run})
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Folder analysis — port of analyzeDocuments(), extended to N documents
# ─────────────────────────────────────────────────────────────────────────────

def doc_tag(i: int) -> str:
    """A, B, …, Z, AA, AB, … (spreadsheet-style)."""
    tag = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        tag = chr(65 + r) + tag
    return tag


def find_docx(folder: Path, recursive: bool) -> list[Path]:
    pattern = "**/*.docx" if recursive else "*.docx"
    files = [p for p in folder.glob(pattern)
             if p.is_file() and not p.name.startswith("~$")]      # skip Word lock files
    return sorted(files, key=lambda p: str(p.relative_to(folder)).lower())


def analyze_folder(folder: Path, recursive: bool = False, log=print,
                   title10: dict | None = None, ecfr_info: dict | None = None) -> dict:
    started = time.time()
    documents: list[dict] = []
    parsed: list[tuple[dict, dict, list]] = []    # (doc, units, instructions)

    files = find_docx(folder, recursive)
    for i, path in enumerate(files):
        rel = str(path.relative_to(folder))
        doc = {"id": i, "tag": doc_tag(i), "name": path.name, "relPath": rel,
               "bytes": path.stat().st_size, "error": None}
        try:
            paras = read_docx_paragraphs(path)
            units, instructions = parse_amendment_doc(paras, title10)
            doc.update({
                "paragraphs": len(paras),
                "sections": len({u["section"] for u in units.values()}),
                "units": len(units),
                "instructions": len(instructions),
                "instructionOnly": sum(1 for x in instructions if not x["producedUnits"]),
            })
            parsed.append((doc, units, instructions))
            log(f"  [{doc['tag']}] {rel}: {len(units)} subsections, {len(instructions)} instructions")
        except Exception as e:                      # one bad file must not stop the run
            doc["error"] = str(e)
            log(f"  [{doc['tag']}] {rel}: ERROR {e}")
        documents.append(doc)

    # ── cross-document index, in document then document-order ──
    index: dict[str, list[tuple[dict, dict, list]]] = {}
    for doc, units, instructions in parsed:
        for key, unit in units.items():
            index.setdefault(key, []).append((doc, unit, instructions))

    conflicts: list[dict] = []
    unique_to_one = agreeing = 0

    for key, entries in index.items():
        if len(entries) < 2:
            unique_to_one += 1
            continue

        # Distinct variants, in order of first appearance.
        variant_of: dict[str, int] = {}
        variants: list[dict] = []
        entry_variant: list[int] = []
        for doc, unit, _ in entries:
            n = substantive_key(unit["text"])         # formatting-only differences: same variant
            if n not in variant_of:
                variant_of[n] = len(variants)
                variants.append({"id": len(variants), "kind": unit["kind"], "text": unit["text"],
                                 "docs": [], "hasElision": False, "_norm": n})
            v = variants[variant_of[n]]
            v["docs"].append(doc["id"])
            v["hasElision"] = v["hasElision"] or unit["hasElision"]
            entry_variant.append(v["id"])

        forms_differ = len({u["kind"] for _, u, _ in entries}) > 1
        if len(variants) < 2:
            agreeing += 1
            continue

        # Most divergent same-kind pair, scanned over entries exactly as the JS does.
        worst = None
        for i in range(len(entries)):
            for j in range(i + 1, len(entries)):
                ui, uj = entries[i][1], entries[j][1]
                if ui["kind"] != uj["kind"]:
                    continue
                if entry_variant[i] == entry_variant[j]:
                    continue
                sim = similarity(ui["text"], uj["text"])
                if worst is None or sim < worst[2]:
                    worst = (i, j, sim)

        form_only = worst is None
        if form_only:
            ti = next((k for k, e in enumerate(entries) if e[1]["kind"] == "text"), 0)
            ii = next((k for k, e in enumerate(entries) if e[1]["kind"] == "instruction"), 1)
            default_pair = [entry_variant[ti], entry_variant[ii]]
            min_sim = None
        else:
            default_pair = [entry_variant[worst[0]], entry_variant[worst[1]]]
            min_sim = worst[2]

        # Redline every same-kind variant pair (variants, not documents, so a
        # text shared by ten documents is diffed once).
        pairs = []
        for i in range(len(variants)):
            for j in range(i + 1, len(variants)):
                va, vb = variants[i], variants[j]
                if va["kind"] != vb["kind"]:
                    continue
                pairs.append({
                    "a": i, "b": j,
                    "similarity": round(similarity(va["text"], vb["text"]), 4),
                    "rows": diff_paragraph_blocks(va["text"], vb["text"]),
                })

        section = entries[0][1]["section"]
        instr = []
        for doc, _, instructions in entries:
            for ins in instructions:
                if section in ins["refs"]:
                    instr.append({"doc": doc["id"], "num": ins["num"], "text": ins["text"]})

        for v in variants:
            del v["_norm"]

        conflicts.append({
            "key": key,
            "section": section,
            "path": entries[0][1]["path"],
            "pathLabel": entries[0][1].get("pathLabel") or entries[0][1]["path"],
            "title": next((u["sectionTitle"] for _, u, _ in entries if u["sectionTitle"]), ""),
            "docs": [d["id"] for d, _, _ in entries],
            "variantCount": len(variants),
            "formsDiffer": forms_differ,
            "formOnly": form_only,
            "hasElision": any(u["hasElision"] for _, u, _ in entries),
            "minSimilarity": None if min_sim is None else round(min_sim, 4),
            "defaultPair": default_pair,
            "variants": variants,
            "pairs": pairs,
            "instructions": instr,
        })

    conflicts.sort(key=lambda c: (_section_sort_key(c["section"]), c["path"]))

    # One instruction aimed at several paragraphs ("in paragraphs (a) and (c),
    # remove …") yields the identical comparison at each; list it once,
    # labelled with every paragraph. Port of the merge in analyzeDocuments().
    merged: list[dict] = []
    by_signature: dict[tuple, dict] = {}
    for c in conflicts:
        if all(v["kind"] == "instruction" for v in c["variants"]):
            sig = (c["section"], tuple(sorted((doc, substantive_key(v["text"]))
                                              for v in c["variants"] for doc in v["docs"])))
            first = by_signature.get(sig)
            if first:
                first["pathLabel"] = f"{first['pathLabel'] or 'section body'}, {c['pathLabel'] or 'section body'}"
                continue
            by_signature[sig] = c
        merged.append(c)
    conflicts = merged
    for i, c in enumerate(conflicts):
        c["id"] = i

    conflict_count: dict[int, int] = {}
    for c in conflicts:
        for d in set(c["docs"]):
            conflict_count[d] = conflict_count.get(d, 0) + 1
    for d in documents:
        d["conflicts"] = conflict_count.get(d["id"], 0)

    return {
        "schema": SCHEMA,
        "generatedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        "tool": "folder_compare.py",
        "folder": str(folder.resolve()),
        "recursive": recursive,
        "ecfr": ecfr_info,            # Title 10 section list used, or None
        "elapsedSeconds": round(time.time() - started, 2),
        "documents": documents,
        "stats": {
            "documents": len(documents),
            "readable": len(parsed),
            "failed": len(documents) - len(parsed),
            "shared": len(index) - unique_to_one,
            "conflicts": len(conflicts),
            "sectionsAffected": len({c["section"] for c in conflicts}),
            "agreeing": agreeing,
            "uniqueToOne": unique_to_one,
            "formOnly": sum(1 for c in conflicts if c["formOnly"]),
        },
        "conflicts": conflicts,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Tabular export — one row per redline
# ─────────────────────────────────────────────────────────────────────────────

def _tags(dataset: dict, ids: list[int]) -> str:
    by_id = {d["id"]: d for d in dataset["documents"]}
    return ", ".join(f"{by_id[i]['tag']} ({by_id[i]['name']})" for i in ids)


def redline_rows(dataset: dict):
    """Yield one flat record per redline: every same-kind variant pair, plus
    one record per text-vs-instruction pairing (shown side by side, not diffed)."""
    for c in dataset["conflicts"]:
        instr = " | ".join(f"{doc_tag(x['doc'])}: {x['text']}" for x in c["instructions"])
        base = {
            "conflict_id": c["id"], "section": c["section"],
            "subsection": c.get("pathLabel") or c["path"] or "(section body)",
            "section_title": c["title"], "variant_count": c["variantCount"],
            "elision": "yes" if c["hasElision"] else "", "instructions": instr,
        }
        vs = c["variants"]
        for p in c["pairs"]:
            va, vb = vs[p["a"]], vs[p["b"]]
            yield {**base, "comparison": va["kind"],
                   "variant_a": p["a"] + 1, "docs_a": _tags(dataset, va["docs"]),
                   "variant_b": p["b"] + 1, "docs_b": _tags(dataset, vb["docs"]),
                   "similarity": p["similarity"], "rows": p["rows"],
                   "text_a": va["text"], "text_b": vb["text"]}
        for i, va in enumerate(vs):
            for j, vb in enumerate(vs):
                if i < j and va["kind"] != vb["kind"]:
                    yield {**base, "comparison": "form differs",
                           "variant_a": i + 1, "docs_a": _tags(dataset, va["docs"]),
                           "variant_b": j + 1, "docs_b": _tags(dataset, vb["docs"]),
                           "similarity": None, "rows": None,
                           "text_a": va["text"], "text_b": vb["text"]}


def rows_to_markup(rows) -> str:
    if rows is None:
        return "(different forms — amendatory text vs. instruction; not redlined)"
    parts = []
    for r in rows:
        if r["t"] == "omit":
            parts.append(f"[… {r['n']} unchanged paragraph{'s' if r['n'] != 1 else ''} …]")
        else:
            parts.append("".join(t if op == "=" else f"[-{t}-]" if op == "-" else f"{{+{t}+}}"
                                 for op, t in r["ops"]))
    return "\n".join(parts)


CSV_COLUMNS = ["conflict_id", "section", "subsection", "section_title", "comparison",
               "variant_a", "docs_a", "variant_b", "docs_b", "similarity", "variant_count",
               "elision", "redline", "text_a", "text_b", "instructions"]


def write_csv(dataset: dict, path: Path) -> int:
    n = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as f:      # BOM so Excel reads UTF-8
        w = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        w.writeheader()
        for r in redline_rows(dataset):
            w.writerow({**{k: r.get(k, "") for k in CSV_COLUMNS},
                        "similarity": "" if r["similarity"] is None else f"{r['similarity']:.2f}",
                        "redline": rows_to_markup(r["rows"])})
            n += 1
    return n


def write_xlsx(dataset: dict, path: Path) -> int:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Font, PatternFill

    LIMIT = 32000                           # Excel caps a cell at 32,767 chars
    clean = lambda s: ILLEGAL_CHARACTERS_RE.sub("", s or "")
    f_del = InlineFont(strike=True, color="FFA01A1A")
    f_ins = InlineFont(u="single", color="FF1A7A3A")
    f_omit = InlineFont(i=True, color="FF888888")

    def rich(rows):
        if rows is None:
            return rows_to_markup(None)
        blocks, used = [], 0

        def add(text, font=None):
            nonlocal used
            text = clean(text)
            if used >= LIMIT or not text:
                return
            text = text[:LIMIT - used]
            used += len(text)
            blocks.append(TextBlock(font, text) if font else text)

        for k, r in enumerate(rows):
            if k:
                add("\n")
            if r["t"] == "omit":
                add(f"[… {r['n']} unchanged paragraph{'s' if r['n'] != 1 else ''} …]", f_omit)
            else:
                for op, t in r["ops"]:
                    add(t, None if op == "=" else f_del if op == "-" else f_ins)
        if used >= LIMIT:
            blocks.append(TextBlock(f_omit, " … [truncated — full redline in the .json]"))
        return CellRichText(blocks)

    wb = Workbook()
    ws = wb.active
    ws.title = "Conflicts"
    header = ["ID", "Section", "Subsection", "Section title", "Comparison",
              "Variant A", "Documents A", "Variant B", "Documents B", "Similarity",
              "Variants", "Elision", "Redline (A → B)", "Text A", "Text B", "Instructions"]
    ws.append(header)
    n = 0
    for r in redline_rows(dataset):
        ws.append([r["conflict_id"], r["section"], r["subsection"], clean(r["section_title"]),
                   r["comparison"], r["variant_a"], clean(r["docs_a"]), r["variant_b"],
                   clean(r["docs_b"]), r["similarity"], r["variant_count"], r["elision"],
                   None, clean(r["text_a"])[:LIMIT], clean(r["text_b"])[:LIMIT],
                   clean(r["instructions"])[:LIMIT]])
        ws.cell(row=ws.max_row, column=13).value = rich(r["rows"])
        n += 1

    widths = [6, 10, 14, 30, 13, 9, 28, 9, 28, 10, 8, 8, 90, 50, 50, 50]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt
    head_fill = PatternFill("solid", fgColor="FF1A1A2E")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = head_fill
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
        row[9].number_format = "0%"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ds = wb.create_sheet("Documents")
    ds.append(["Tag", "File", "Subsections", "Sections", "Instructions",
               "Instruction-only", "Conflicts", "Error"])
    for d in dataset["documents"]:
        ds.append([d["tag"], d["relPath"], d.get("units"), d.get("sections"), d.get("instructions"),
                   d.get("instructionOnly"), d.get("conflicts"), d["error"] or ""])
    for i, wdt in enumerate([6, 50, 12, 10, 12, 16, 10, 50], start=1):
        ds.column_dimensions[ds.cell(row=1, column=i).column_letter].width = wdt
    for c in ds[1]:
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = head_fill

    ss = wb.create_sheet("Summary")
    ecfr = dataset.get("ecfr")
    ecfr_text = (f"Title 10 as of {ecfr['date']} ({ecfr['source']})" if ecfr
                 else "not used — unmarked numbers counted as sections only when an instruction named them")
    for k, v in [("Folder", dataset["folder"]), ("Generated", dataset["generatedAt"]),
                 ("eCFR section list", ecfr_text),
                 *[(k, v) for k, v in dataset["stats"].items()]]:
        ss.append([k, v])
    ss.column_dimensions["A"].width = 20
    ss.column_dimensions["B"].width = 80

    wb.save(path)
    return n


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Compare every .docx amendment document in a folder and write a "
                    "dataset of conflicting subsections with redlines.")
    ap.add_argument("folder", type=Path, help="folder containing .docx files")
    ap.add_argument("-o", "--out-dir", type=Path, help="where to write outputs (default: the folder)")
    ap.add_argument("--name", default="regdocdiff-folder", help="output file base name")
    ap.add_argument("-r", "--recursive", action="store_true", help="include subfolders")
    ap.add_argument("--no-csv", action="store_true", help="skip the .csv output")
    ap.add_argument("--no-xlsx", action="store_true", help="skip the .xlsx output")
    ap.add_argument("--no-ecfr", action="store_true",
                    help="don't check section numbers against the eCFR Title 10 list (then an "
                         "unmarked number counts as a section only when an instruction names it)")
    ap.add_argument("--ecfr-cache", type=Path,
                    help=f"where to cache the eCFR section list (default: {default_cache_dir()})")
    ap.add_argument("-q", "--quiet", action="store_true", help="only print output paths")
    args = ap.parse_args(argv)

    folder = args.folder
    if not folder.is_dir():
        print(f"error: {folder} is not a folder", file=sys.stderr)
        return 2
    log = (lambda *_: None) if args.quiet else print
    warn = lambda msg: print(msg, file=sys.stderr)      # warnings show even with --quiet

    title10, ecfr_info = None, None
    if not args.no_ecfr:
        log("Loading the eCFR Title 10 section list…")
        title10, ecfr_info = load_title10_sections(args.ecfr_cache, warn)
        if ecfr_info:
            log(f"  {ecfr_info['sections']} sections as of {ecfr_info['date']} ({ecfr_info['source']})")

    log(f"Scanning {folder.resolve()}{' (recursive)' if args.recursive else ''}")
    dataset = analyze_folder(folder, args.recursive, log, title10, ecfr_info)
    s = dataset["stats"]
    if s["documents"] == 0:
        print("error: no .docx files found", file=sys.stderr)
        return 1
    log(f"\n{s['readable']} readable document(s), {s['failed']} failed · "
        f"{s['shared']} shared subsection(s) · {s['conflicts']} conflict(s) in "
        f"{s['sectionsAffected']} section(s) · {s['agreeing']} in agreement · "
        f"{dataset['elapsedSeconds']}s")
    if s["readable"] < 2:
        log("note: fewer than two readable documents — nothing to compare")

    out = args.out_dir or folder
    out.mkdir(parents=True, exist_ok=True)

    json_path = out / f"{args.name}.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(dataset, f, ensure_ascii=False, separators=(",", ":"))
    print(f"wrote {json_path}")

    if not args.no_csv:
        p = out / f"{args.name}.csv"
        n = write_csv(dataset, p)
        print(f"wrote {p} ({n} rows)")

    if not args.no_xlsx:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            log("skipped .xlsx (openpyxl not installed: pip install openpyxl)")
        else:
            p = out / f"{args.name}.xlsx"
            try:
                n = write_xlsx(dataset, p)
                print(f"wrote {p} ({n} rows)")
            except PermissionError:
                print(f"error: could not write {p} — is it open in Excel?", file=sys.stderr)
                return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
